#!/usr/bin/env python3
"""Build docs/pitch-deck-v3.pptx — Ember, reframed for Stellar.
Base/reference: docs/pitch-deck/ember-pitch-deck.md (brand + structure).
Content refs: GitHub issue #179 (Stellar ecosystem strategy) + #177 (migration)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Brand palette (from ember-pitch-deck.md) ----
EMBER      = RGBColor(0xB7, 0x23, 0x01)
EMBER_BR   = RGBColor(0xFF, 0x57, 0x33)
EMBER_SOFT = RGBColor(0xFF, 0xDA, 0xD3)
GREEN      = RGBColor(0x00, 0x6C, 0x4B)
GREEN_BR   = RGBColor(0x00, 0xA5, 0x74)
GREEN_SOFT = RGBColor(0xC8, 0xF7, 0xE4)
INK        = RGBColor(0x19, 0x1C, 0x1D)
MUTED      = RGBColor(0x5B, 0x40, 0x3A)
BG         = RGBColor(0xF8, 0xF9, 0xFA)
SURFACE    = RGBColor(0xFF, 0xFF, 0xFF)
LINE       = RGBColor(0xE1, 0xE3, 0xE4)
CARDLINE   = RGBColor(0xE4, 0xBE, 0xB6)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
CREAM      = RGBColor(0xFF, 0xD9, 0xD0)
DARK       = RGBColor(0x19, 0x1C, 0x1D)
DARKRED    = RGBColor(0x2A, 0x0C, 0x05)
CHIP_BG    = RGBColor(0xED, 0xEE, 0xEF)

FONT = "Inter"
MONO = "JetBrains Mono"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def _set_run(run, text, size, color, bold=False, font=FONT, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def textbox(s, left, top, width, height, runs, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, space_after=6, line_spacing=1.12):
    """runs: list of paragraphs; each paragraph is list of (text,size,color,bold,font) tuples."""
    tb = s.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for seg in para:
            text, size, color = seg[0], seg[1], seg[2]
            bold = seg[3] if len(seg) > 3 else False
            font = seg[4] if len(seg) > 4 else FONT
            italic = seg[5] if len(seg) > 5 else False
            _set_run(p.add_run(), text, size, color, bold, font, italic)
    return tb


def card(s, left, top, width, height, fill=SURFACE, line=LINE, radius=0.10):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.color.rgb = line; c.line.width = Pt(1)
    c.shadow.inherit = False
    try:
        c.adjustments[0] = radius
    except Exception:
        pass
    return c


def pill(s, left, top, text, fill=GREEN_SOFT, color=RGBColor(0x00,0x31,0x20), width=Inches(2.0)):
    h = Inches(0.36)
    p = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, h)
    p.fill.solid(); p.fill.fore_color.rgb = fill
    p.line.fill.background(); p.shadow.inherit = False
    try: p.adjustments[0] = 0.5
    except Exception: pass
    tf = p.text_frame; tf.word_wrap = False
    tf.margin_top = 0; tf.margin_bottom = 0
    par = tf.paragraphs[0]; par.alignment = PP_ALIGN.CENTER
    _set_run(par.add_run(), text, 12, color, True)
    return p


def eyebrow(s, left, top, text, color=EMBER, width=Inches(9)):
    textbox(s, left, top, width, Inches(0.3),
            [[(text.upper(), 12.5, color, True)]])


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


MX = Inches(0.75)   # left margin
CW = Inches(11.83)  # content width


def title(s, text, top=Inches(1.02), color=INK, size=33, width=CW):
    textbox(s, MX, top, width, Inches(1.0), [[(text, size, color, True)]])


def footer(s, text="Ember · Fund the future, one milestone at a time", dark=False):
    col = RGBColor(0x8C,0x7A,0x75) if dark else MUTED
    textbox(s, MX, Inches(7.08), CW, Inches(0.3),
            [[(text, 10.5, col)]])


# =====================================================================
# SLIDE 1 — COVER
# =====================================================================
s = slide(DARK)
# radial-ish: dark red panel top-right
g = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
g.fill.solid(); g.fill.fore_color.rgb = DARK
g.line.fill.background(); g.shadow.inherit = False
# ember accent panel (top-right corner glow, solid)
acc = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.2), Inches(-3.2), Inches(9.0), Inches(7.0))
acc.fill.solid(); acc.fill.fore_color.rgb = DARKRED
acc.line.fill.background(); acc.shadow.inherit = False

textbox(s, MX, Inches(1.5), CW, Inches(0.4), [[("MILESTONE-BASED ONCHAIN CROWDFUNDING", 13, RGBColor(0xC8,0xB7,0xB2), True)]])
textbox(s, MX, Inches(2.1), CW, Inches(1.6),
        [[("Ember", 82, WHITE, True), (".", 82, EMBER_BR, True)]])
textbox(s, MX, Inches(3.65), CW, Inches(0.7),
        [[("Fund the future, one milestone at a time.", 30, CREAM, False)]])
textbox(s, MX, Inches(4.7), CW, Inches(0.5),
        [[("Trust-minimized crowdfunding · Migrating to ", 17, RGBColor(0xC8,0xB7,0xB2)),
          ("Stellar / Soroban", 17, EMBER_BR, True)]])
pill(s, MX, Inches(5.5), "PITCH DECK v3 — STELLAR EDITION", fill=EMBER, color=WHITE, width=Inches(3.6))
footer(s, "Ember · 2026 · references GitHub issue #179 (Stellar ecosystem strategy)", dark=True)
notes(s, "Opening line: 'Ember is trust-minimized crowdfunding — funds release as creators "
        "deliver, not before.' This is v3 of the deck, reframed around our move to Stellar. "
        "The product is already built and feature-complete on an EVM testnet today; this deck "
        "makes the case for why Stellar is the right home and how we drive real, measurable "
        "impact for the ecosystem. Set the tone: a working product with a clear ecosystem thesis "
        "and a concrete go-to-market, not a whitepaper. Everything technical in this deck traces "
        "back to the research report saved as issue #179.")

# =====================================================================
# SLIDE 2 — PROBLEM
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "The problem")
title(s, "Crowdfunding is broken — and backers pay for it")
textbox(s, MX, Inches(1.85), CW, Inches(1.0),
        [[("Platforms like Kickstarter hand creators ", 21, INK), ("100% of the funds upfront", 21, EMBER, True),
          (" — before a single promise is kept. When a project stalls, ", 21, INK),
          ("the backer eats the loss.", 21, EMBER, True), (" No recourse, no refund, no leverage.", 21, INK)]],
        line_spacing=1.3)
cw = Inches(5.75); ch = Inches(2.7); cy = Inches(3.5)
card(s, MX, cy, cw, ch)
textbox(s, MX+Inches(0.35), cy+Inches(0.3), cw-Inches(0.7), Inches(0.5), [[("Today", 22, INK, True)]])
textbox(s, MX+Inches(0.35), cy+Inches(0.95), cw-Inches(0.7), Inches(1.6),
        [[("•  Funds released on day one", 18, INK)],
         [("•  Trust sits with a platform", 18, INK)],
         [("•  Disputes settled by support tickets", 18, INK)],
         [("•  High fees on every transaction", 18, INK)]], line_spacing=1.3, space_after=8)
c2x = MX + cw + Inches(0.33)
card(s, c2x, cy, cw, ch)
textbox(s, c2x+Inches(0.35), cy+Inches(0.3), cw-Inches(0.7), Inches(0.5), [[("The cost", 22, EMBER, True)]])
textbox(s, c2x+Inches(0.35), cy+Inches(0.95), cw-Inches(0.7), Inches(1.6),
        [[("•  Broken promises, no accountability", 18, INK)],
         [("•  Capital locked away from those who need it", 18, INK)],
         [("•  Worst in emerging markets that need it most", 18, INK)],
         [("•  No portable, verifiable proof of contribution", 18, INK)]], line_spacing=1.3, space_after=8)
footer(s)
notes(s, "Anchor the pain: on Web2 crowdfunding, money moves to the creator on day one and trust "
        "is just a promise. When delivery fails, the backer has no leverage — only a support ticket. "
        "Emphasize that this failure is worst exactly where accountable capital matters most: "
        "emerging markets and small creators. Keep it short — the audience already feels this. "
        "The point of this slide is to set up that the fix is structural, not customer-service.")

# =====================================================================
# SLIDE 3 — WHY SEA / PH
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Why here, why now")
title(s, "In Southeast Asia, the gap is widest")
lx = MX; lw = Inches(6.5)
textbox(s, lx, Inches(1.95), lw, Inches(4.0),
        [[("•  The ", 19, INK), ("Philippines", 19, EMBER, True), (" is among the world's leaders in crypto-wallet adoption.", 19, INK)],
         [("•  Millions already hold and move ", 19, INK), ("stablecoins", 19, EMBER, True), (" for remittances and savings.", 19, INK)],
         [("•  A fast-growing ", 19, INK), ("creator economy", 19, EMBER, True), (" and ", 19, INK), ("SME", 19, EMBER, True), (" base — underserved by banks.", 19, INK)],
         [("•  ", 19, INK), ("$30B+", 19, EMBER, True), (" in annual remittances flow into the country every year.", 19, INK)],
         [("•  Small businesses lack the escrow, controls, and capital access big institutions take for granted.", 19, INK)]],
        line_spacing=1.28, space_after=14)
rx = MX + lw + Inches(0.4); rw = Inches(4.93)
card(s, rx, Inches(1.95), rw, Inches(3.7), fill=EMBER, line=EMBER)
textbox(s, rx+Inches(0.4), Inches(2.35), rw-Inches(0.8), Inches(2.2),
        [[("The infrastructure to fix this already exists onchain.", 27, WHITE, True)]], line_spacing=1.2)
textbox(s, rx+Inches(0.4), Inches(4.35), rw-Inches(0.8), Inches(1.2),
        [[("What's missing is a product that puts it in the hands of everyday creators and small businesses — and a chain built for stablecoin payments.", 16.5, CREAM)]], line_spacing=1.25)
footer(s)
notes(s, "This is the 'why us, why now, why this region' slide. The Philippines is a global leader in "
        "crypto adoption with $30B+ in annual remittances and a population already fluent in moving "
        "stablecoins — but SMEs and creators are locked out of formal capital and escrow. Land the "
        "punchline in the red card: the onchain infrastructure exists; what's missing is a product "
        "that delivers it to everyday users, on a chain purpose-built for stablecoin payments. That "
        "chain is Stellar — which is the natural segue into the solution and the ecosystem thesis.")

# =====================================================================
# SLIDE 4 — DIVIDER: EMBER
# =====================================================================
s = slide(EMBER)
textbox(s, MX, Inches(2.55), CW, Inches(1.4), [[("Ember", 60, WHITE, True)]])
textbox(s, MX, Inches(3.75), Inches(10.5), Inches(1.6),
        [[("Onchain crowdfunding where funds release ", 26, CREAM),
          ("as creators deliver — not before.", 26, WHITE, True)]], line_spacing=1.25)
footer(s, "Ember", dark=False)
notes(s, "Breath slide. Say the one-sentence definition out loud and pause: 'Ember is onchain "
        "crowdfunding where funds release as creators deliver — not before.' Then move into the "
        "mechanics on the next slide.")

# =====================================================================
# SLIDE 5 — WHAT EMBER DOES
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "The solution")
title(s, "What Ember does")
textbox(s, MX, Inches(1.85), CW, Inches(1.0),
        [[("Backers contribute a ", 20, INK), ("stablecoin", 20, EMBER, True),
          (" into a project's own ", 20, INK), ("escrow smart contract", 20, EMBER, True),
          (" and receive a ", 20, INK), ("position NFT", 20, EMBER, True),
          (" as proof of stake. Funds release to the creator ", 20, INK),
          ("only when backers vote to approve each milestone.", 20, EMBER, True)]], line_spacing=1.3)
cy = Inches(3.55); cw = Inches(3.77); ch = Inches(2.75); gap = Inches(0.26)
cards = [
    ("Trust-first", "Cryptographic guarantees, not platform promises. Non-custodial — Ember never holds your funds.", EMBER),
    ("Algorithmic", "Code decides fund release. No human discretion, no customer-service-as-justice.", EMBER),
    ("Stablecoin-native", "USDC on Stellar. No volatile tokens on a pledge. Sub-cent fees, ~5s settlement.", GREEN),
]
for i, (h, b, col) in enumerate(cards):
    x = MX + i*(cw+gap)
    card(s, x, cy, cw, ch)
    textbox(s, x+Inches(0.32), cy+Inches(0.32), cw-Inches(0.64), Inches(0.5), [[(h, 20, col, True)]])
    textbox(s, x+Inches(0.32), cy+Inches(0.95), cw-Inches(0.64), Inches(1.7), [[(b, 17, INK)]], line_spacing=1.3)
footer(s)
notes(s, "Three pillars. Trust-first: it's non-custodial — Ember never has a key that can move backer "
        "funds; the contract does. Algorithmic: code releases funds on a weighted vote, no human "
        "discretion. Stablecoin-native: note the deliberate shift here — we now say USDC on Stellar, "
        "not USDT on Morph. That is the core of v3: USDC is Stellar's natively-supported settlement "
        "stablecoin, with sub-cent fees and ~5-second finality. If asked why we changed the asset, "
        "the short answer is on issue #179 — USDC is the dominant, native Stellar stablecoin.")

# =====================================================================
# SLIDE 6 — HOW IT WORKS
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Four steps, fully onchain")
title(s, "How it works")
rows = [
    ("1 · Create", "A verified org publishes a project with milestones", "Factory deploys a dedicated escrow + position NFT"),
    ("2 · Back", "Backers contribute stablecoin, mint a position NFT", "Escrow.contribute() — funds held onchain"),
    ("3 · Vote", "Each milestone is approved by weighted backer vote", "vote() → resolveMilestone()"),
    ("4 · Release", "Funds release only for passed milestones", "claimMilestone() → USDC to creator"),
]
ty = Inches(1.95); rowh = Inches(0.92)
c1 = Inches(2.4); c2 = Inches(5.7); c3 = Inches(3.73)
# header
hx = MX
card(s, MX, ty, CW, Inches(0.55), fill=CHIP_BG, line=CHIP_BG, radius=0.3)
textbox(s, MX+Inches(0.25), ty+Inches(0.11), c1, Inches(0.4), [[("Step", 15, INK, True)]])
textbox(s, MX+Inches(0.25)+c1, ty+Inches(0.11), c2, Inches(0.4), [[("What happens", 15, INK, True)]])
textbox(s, MX+Inches(0.25)+c1+c2, ty+Inches(0.11), c3, Inches(0.4), [[("Onchain", 15, INK, True)]])
for i, (a, b, c) in enumerate(rows):
    y = ty + Inches(0.62) + i*rowh
    textbox(s, MX+Inches(0.25), y+Inches(0.16), c1, Inches(0.7), [[(a, 17, EMBER, True)]])
    textbox(s, MX+Inches(0.25)+c1, y+Inches(0.16), c2-Inches(0.2), Inches(0.7), [[(b, 15.5, INK)]], line_spacing=1.15)
    textbox(s, MX+Inches(0.25)+c1+c2, y+Inches(0.16), c3, Inches(0.7), [[(c, 13.5, MUTED, False, MONO)]], line_spacing=1.15)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MX, y+rowh-Emu(9525), CW, Emu(9525))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); ln.shadow.inherit=False
textbox(s, MX, Inches(6.35), CW, Inches(0.6),
        [[("A permissionless ", 17, INK), ("keeper", 17, EMBER, True),
          (" resolves milestones when the voting window closes — no one can stall a payout.", 17, INK)]])
footer(s)
notes(s, "Walk the four steps as one continuous flow. Emphasize step 4: money only ever leaves escrow "
        "for a milestone backers approved. Call out the keeper: it's permissionless, so resolution "
        "and payout can't be censored or stalled by us or the creator. On Stellar this maps cleanly — "
        "the escrow becomes a Soroban contract, the position NFT uses the SEP-50 standard, and the "
        "keeper role can even be simplified. Note the function names shown are our current contract "
        "API; the Soroban port keeps the same shape.")

# =====================================================================
# SLIDE 7 — BUILT AND WORKING
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "What's live today")
title(s, "A working product, not a whitepaper")
lx = MX; lw = Inches(6.6)
textbox(s, lx, Inches(1.95), lw, Inches(4.2),
        [[("•  ", 18, INK), ("Feature-complete", 18, EMBER, True), (" against the full product spec — end to end.", 18, INK)],
         [("•  Create → publish → back → vote → release, fully onchain.", 18, INK)],
         [("•  Backer & org ", 18, INK), ("dashboards", 18, EMBER, True), (", in-app + email notifications.", 18, INK)],
         [("•  ", 18, INK), ("Admin panel", 18, EMBER, True), (": org verification, audit logs, reports.", 18, INK)],
         [("•  Standalone ", 18, INK), ("event indexer", 18, EMBER, True), (", permissionless keeper, immutable contracts.", 18, INK)],
         [("•  Foundry test suite: unit, fuzz & invariant coverage.", 18, INK)]],
        line_spacing=1.3, space_after=12)
rx = MX + lw + Inches(0.35); rw = Inches(4.88)
card(s, rx, Inches(1.95), rw, Inches(4.15))
pill(s, rx+Inches(0.35), Inches(2.3), "BUILT ON EVM TESTNET TODAY", fill=GREEN_SOFT, color=RGBColor(0x00,0x31,0x20), width=Inches(3.4))
textbox(s, rx+Inches(0.35), Inches(2.95), rw-Inches(0.7), Inches(1.3),
        [[("The hard part is done.", 26, INK, True)]], line_spacing=1.1)
textbox(s, rx+Inches(0.35), Inches(3.95), rw-Inches(0.7), Inches(2.0),
        [[("Web app, indexer, and immutable contracts run together as one system. ", 16.5, INK),
          ("Stellar is a chain swap on a proven product", 16.5, EMBER, True),
          (" — not a rebuild.", 16.5, INK)]], line_spacing=1.3)
footer(s)
notes(s, "De-risk the bet. The product is not a concept — it's feature-complete against our full spec: "
        "the entire create→back→vote→release loop, dashboards, admin, indexer, keeper, and a Foundry "
        "test suite with fuzz and invariant tests, all deployed to testnet. The key investor message: "
        "moving to Stellar is a chain-integration swap on a working product, not a from-scratch build. "
        "That massively lowers execution risk. If pressed on current state: yes, it runs on an EVM "
        "testnet today; the Soroban migration is scoped in issue #177 and the ecosystem plan in #179.")

# =====================================================================
# SLIDE 8 — IMPACT TO STELLAR ECOSYSTEM  (NEW)
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Why Stellar — and what we bring to it")
title(s, "Impact on the Stellar ecosystem")
textbox(s, MX, Inches(1.75), CW, Inches(0.7),
        [[("Ember is an engine for the metrics Stellar cares about most — turning real-world giving and SME capital into onchain activity.", 18, INK)]], line_spacing=1.25)
stats = [
    ("Stablecoin\nvolume", "Every contribution & milestone payout settles in USDC onchain", EMBER),
    ("Wallet\nactivations", "Each passkey onboard = a net-new activated Stellar wallet", GREEN),
    ("Real-world\nimpact", "Remittance & diaspora-giving flows become measurable TVL", EMBER),
    ("New\ncategory", "Accountable milestone crowdfunding Stellar doesn't yet have", GREEN),
]
cy = Inches(2.75); cw = Inches(2.86); ch = Inches(2.55); gap = Inches(0.17)
for i, (h, b, col) in enumerate(stats):
    x = MX + i*(cw+gap)
    card(s, x, cy, cw, ch)
    textbox(s, x+Inches(0.28), cy+Inches(0.3), cw-Inches(0.56), Inches(1.0), [[(h, 21, col, True)]], line_spacing=1.05)
    textbox(s, x+Inches(0.28), cy+Inches(1.45), cw-Inches(0.56), Inches(1.0), [[(b, 15, INK)]], line_spacing=1.25)
card(s, MX, Inches(5.6), CW, Inches(1.15), fill=EMBER_SOFT, line=CARDLINE)
textbox(s, MX+Inches(0.4), Inches(5.82), CW-Inches(0.8), Inches(0.8),
        [[("A strong ", 17, RGBColor(0x8C,0x18,0x00)), ("Stellar Community Fund", 17, EMBER, True),
          (" candidate: a Soroban dApp driving stablecoin settlement, wallet growth, and a real “on-chain impact” use case — precisely SCF 7.0's thesis (Build Award up to $150K XLM).", 17, RGBColor(0x8C,0x18,0x00))]], line_spacing=1.25)
footer(s)
notes(s, "This is the ecosystem-impact slide — the reason a Stellar audience should care. Four levers: "
        "(1) stablecoin volume — every contribution and payout is USDC settlement; (2) wallet "
        "activations — passkey onboarding creates net-new Stellar accounts from non-crypto backers; "
        "(3) real-world impact/TVL — we convert remittance and diaspora-giving flows into onchain "
        "activity; (4) new category — accountable, milestone-gated crowdfunding doesn't exist on "
        "Stellar yet. Close on the SCF angle: this maps exactly to SCF 7.0's 'real on-chain impact' "
        "thesis, Build Award up to $150K XLM. All figures and the SCF program detail come from the "
        "issue #179 research report. Flag that SCF cadence/amounts are cited there with verify notes.")

# =====================================================================
# SLIDE 9 — ECOSYSTEM INTEGRATIONS PLAN  (NEW)
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Plans to integrate existing ecosystem solutions")
title(s, "Built on Stellar's rails — not around them")
rows = [
    ("Soroban + SEP-41 / SAC", "Rewrite escrow, voting & milestone logic; hold USDC natively"),
    ("SEP-50 + OpenZeppelin NFT", "Audited position NFTs (ERC-721-like) for backers"),
    ("Anchors — SEP-24 / 6 / 10 / 12 / 38", "Hosted fiat ↔ USDC on/off-ramp, auth, KYC & quotes"),
    ("SEP-31 + Coins.ph (PH anchor)", "Cross-border payouts & a live PHP ramp for OFW backers"),
    ("Passkeys — passkey-kit + Launchtube", "Seedless smart-wallet onboarding + sponsored (gasless) fees"),
    ("Path payments / SDEX / Soroswap", "Backer pays any asset; escrow always accumulates USDC"),
    ("Mercury / Goldsky + Stellar RPC", "Replace our custom indexer with native data infra"),
    ("Stellar Disbursement Platform", "Bulk milestone payouts to many local recipients"),
]
ty = Inches(1.85); rowh = Inches(0.585)
c1 = Inches(5.1)
for i, (a, b) in enumerate(rows):
    y = ty + i*rowh
    if i % 2 == 0:
        bandc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MX, y, CW, rowh)
        bandc.fill.solid(); bandc.fill.fore_color.rgb = SURFACE
        bandc.line.fill.background(); bandc.shadow.inherit=False
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, MX+Inches(0.1), y+Inches(0.22), Inches(0.13), Inches(0.13))
    dot.fill.solid(); dot.fill.fore_color.rgb = EMBER if i%2==0 else GREEN; dot.line.fill.background(); dot.shadow.inherit=False
    textbox(s, MX+Inches(0.4), y+Inches(0.12), c1, Inches(0.5), [[(a, 15.5, INK, True, MONO)]], line_spacing=1.05)
    textbox(s, MX+Inches(0.4)+c1, y+Inches(0.12), CW-c1-Inches(0.5), Inches(0.5), [[(b, 15.5, MUTED)]], line_spacing=1.05)
textbox(s, MX, Inches(6.62), CW, Inches(0.5),
        [[("Full integration matrix, use-cases & reference links: ", 13.5, MUTED),
          ("GitHub issue #179", 13.5, EMBER, True)]])
footer(s)
notes(s, "This is the 'we compose Stellar, we don't reinvent it' slide. Walk top to bottom: core "
        "contracts on Soroban using the SEP-41 token interface so USDC is native; position NFTs on "
        "the SEP-50 standard with OpenZeppelin's audited modules; anchors (the SEP-24 family) give us "
        "a hosted fiat on/off-ramp so we never touch KYC; SEP-31 plus Coins.ph gives a live Philippine "
        "peso ramp; passkeys plus Launchtube deliver seedless, gasless onboarding for non-crypto "
        "backers; path payments let a backer pay any asset while the escrow always ends up in USDC; "
        "Mercury/Goldsky replace our hand-rolled indexer; and the Stellar Disbursement Platform lets "
        "us pay many local recipients at milestone release. Every row has a concrete use-case and "
        "reference URL in issue #179. Message: minimal net-new trusted code, maximum ecosystem reuse.")

# =====================================================================
# SLIDE 10 — GTM PHILIPPINES  (NEW)
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Go-to-market · Level 1")
title(s, "Philippines — win the beachhead")
textbox(s, MX, Inches(1.7), CW, Inches(0.6),
        [[("Land where stablecoin fluency, remittances, and underserved SMEs converge.", 18, INK)]])
cols = [
    ("Wedge", ["Diaspora & OFW-funded community projects and SME raises",
               "Fiat-peso on-ramp via Coins.ph (SEP-24) — back in pesos, settle USDC",
               "Passkey onboarding: no seed phrase, no crypto experience needed"]),
    ("Channels", ["Partner with local creator communities, cooperatives & LGUs",
                  "Stellar PH developer & hackathon community (Stellar Rise PH)",
                  "Crypto-native fintech partners (Coins.ph) for distribution"]),
    ("Proof goals", ["First 25 verified orgs + funded milestone cohort",
                     "Measured: USDC settled, wallets activated, milestone pass-rate",
                     "Case studies to anchor SCF application & APAC expansion"]),
]
cw = Inches(3.77); ch = Inches(3.5); gap = Inches(0.26); cy = Inches(2.45)
for i, (h, items) in enumerate(cols):
    x = MX + i*(cw+gap)
    card(s, x, cy, cw, ch)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, cy, cw, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = EMBER if i!=2 else GREEN; bar.line.fill.background(); bar.shadow.inherit=False
    textbox(s, x+Inches(0.3), cy+Inches(0.35), cw-Inches(0.6), Inches(0.5), [[(h, 20, INK, True)]])
    textbox(s, x+Inches(0.3), cy+Inches(1.0), cw-Inches(0.6), Inches(2.4),
            [[("•  "+it, 15, INK)] for it in items], line_spacing=1.22, space_after=10)
footer(s)
notes(s, "GTM level 1 — the beachhead. The Philippines is deliberately first: it's where stablecoin "
        "fluency, $30B+ remittances, and underserved SMEs overlap. Wedge is diaspora/OFW-funded "
        "community and SME projects, with a peso on-ramp through Coins.ph and passkey onboarding so a "
        "non-crypto tita can back a project in pesos in under a minute. Channels: local creator "
        "communities, cooperatives, LGUs, and the active Stellar PH developer/hackathon scene. Proof "
        "goals are concrete and measurable — first ~25 verified orgs, a funded milestone cohort, and "
        "hard numbers (USDC settled, wallets activated, pass-rate) that both prove the model and "
        "become the SCF application and the springboard to APAC.")

# =====================================================================
# SLIDE 11 — GTM APAC  (NEW)
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Go-to-market · Level 2")
title(s, "APAC — replicate the corridor")
lx = MX; lw = Inches(6.7)
textbox(s, lx, Inches(1.9), lw, Inches(4.5),
        [[("The Philippine playbook is a ", 18, INK), ("template", 18, EMBER, True), (", not a one-off:", 18, INK)],
         [("•  Expand to high-remittance, high-adoption corridors: ", 16.5, INK), ("Vietnam, Indonesia, India", 16.5, EMBER, True), (".", 16.5, INK)],
         [("•  Each market = plug in a local ", 16.5, INK), ("SEP-24 anchor", 16.5, EMBER, True), (" for its fiat currency.", 16.5, INK)],
         [("•  ", 16.5, INK), ("EURC / PYUSD / local stablecoins", 16.5, EMBER, True), (" widen denomination options for backers.", 16.5, INK)],
         [("•  Cross-border SME financing via ", 16.5, INK), ("SEP-31", 16.5, EMBER, True), (" corridors between anchors.", 16.5, INK)],
         [("•  Localized language, curation & compliance per market.", 16.5, INK)]],
        line_spacing=1.28, space_after=12)
rx = MX + lw + Inches(0.35); rw = Inches(4.78)
card(s, rx, Inches(1.9), rw, Inches(4.3), fill=GREEN, line=GREEN)
textbox(s, rx+Inches(0.38), Inches(2.25), rw-Inches(0.76), Inches(1.3),
        [[("One integration, many markets.", 24, WHITE, True)]], line_spacing=1.1)
textbox(s, rx+Inches(0.38), Inches(3.5), rw-Inches(0.76), Inches(2.5),
        [[("Because every market rides the same Stellar rails, expansion cost is an ", 16, RGBColor(0xE6,0xFF,0xF6)),
          ("anchor integration — not a new payments stack.", 16, WHITE, True),
          (" Stellar's global anchor network is the distribution moat.", 16, RGBColor(0xE6,0xFF,0xF6))]], line_spacing=1.3)
footer(s)
notes(s, "GTM level 2 — regional replication. The insight: the PH launch produces a repeatable "
        "corridor playbook. Each new APAC market (Vietnam, Indonesia, India — all high-remittance, "
        "high-adoption) is mostly a matter of plugging in that country's local SEP-24 anchor for its "
        "fiat currency; the product, contracts, and UX are unchanged. EURC, PYUSD and local "
        "stablecoins widen backer denomination choices, and SEP-31 opens cross-border SME financing "
        "between anchors. The green card is the strategic point for investors: expansion cost is an "
        "anchor integration, not a new payments stack — Stellar's global anchor network is our "
        "distribution moat. This is what makes the model scale capital-efficiently.")

# =====================================================================
# SLIDE 12 — GTM GLOBAL  (NEW)
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Go-to-market · Level 3")
title(s, "Global — the accountable-capital layer")
textbox(s, MX, Inches(1.7), CW, Inches(0.6),
        [[("From remittance corridors to a worldwide standard for milestone-gated funding.", 18, INK)]])
tiles = [
    ("Public goods & QF", "Quadratic-funding matching pools on Soroban for open-source & community goods — a novel Stellar primitive."),
    ("RWA & impact funding", "Disaster relief, agri/SME finance, local infrastructure — payouts via the Disbursement Platform."),
    ("Recurring & subscriptions", "Passkey smart wallets enable monthly pledges — smoothing volume, raising lifetime value."),
    ("Position-NFT primitive", "Portable proof-of-contribution with reputation & provenance — a reusable ecosystem building block."),
]
cw = Inches(5.78); ch = Inches(1.75); gapx = Inches(0.27); gapy = Inches(0.22); x0 = MX; y0 = Inches(2.5)
for i, (h, b) in enumerate(tiles):
    x = x0 + (i % 2)*(cw+gapx)
    y = y0 + (i // 2)*(ch+gapy)
    card(s, x, y, cw, ch)
    textbox(s, x+Inches(0.32), y+Inches(0.26), cw-Inches(0.64), Inches(0.5), [[(h, 18.5, EMBER if i%2==0 else GREEN, True)]])
    textbox(s, x+Inches(0.32), y+Inches(0.78), cw-Inches(0.64), Inches(0.9), [[(b, 15, INK)]], line_spacing=1.25)
footer(s)
notes(s, "GTM level 3 — the global vision. Once the corridor model is proven, Ember becomes the "
        "accountable-capital layer for Stellar worldwide, across four expansion vectors: (1) public "
        "goods via quadratic-funding matching pools — which don't exist natively on Stellar yet, so "
        "it's a fundable, novel primitive; (2) real-world-asset and impact funding — disaster relief, "
        "agri/SME finance, infrastructure — with payouts through the Disbursement Platform; (3) "
        "recurring/subscription pledges enabled by passkey smart wallets; (4) the position NFT as a "
        "reusable, portable proof-of-contribution primitive with reputation. Frame this as optionality "
        "on top of a proven base — not promises we need funded today. Each is sourced in issue #179.")

# =====================================================================
# SLIDE 13 — BUSINESS MODEL
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Business model")
title(s, "How Ember makes money")
lx = MX; lw = Inches(6.6)
textbox(s, lx, Inches(1.95), lw, Inches(4.0),
        [[("•  A ", 19, INK), ("small platform fee", 19, EMBER, True), (" on each ", 19, INK), ("milestone release", 19, EMBER, True), (" — we earn only when creators actually deliver.", 19, INK)],
         [("•  Revenue is ", 19, INK), ("aligned with backers", 19, EMBER, True), (": no delivery, no fee.", 19, INK)],
         [("•  Optional ", 19, INK), ("fiat-ramp & FX spread", 19, EMBER, True), (" on anchor-powered on/off-ramps.", 19, INK)],
         [("•  ", 19, INK), ("No token. No sale. No custody.", 19, EMBER, True), (" Pledges are utility, not securities.", 19, INK)]],
        line_spacing=1.3, space_after=14)
rx = MX + lw + Inches(0.35); rw = Inches(4.88)
card(s, rx, Inches(1.95), rw, Inches(3.9))
textbox(s, rx+Inches(0.35), Inches(2.3), rw-Inches(0.7), Inches(0.6), [[("Why it scales", 20, GREEN, True)]])
textbox(s, rx+Inches(0.35), Inches(2.95), rw-Inches(0.7), Inches(2.8),
        [[("Every project deploys its own escrow, and Stellar's ", 16.5, INK),
          ("sub-cent fees", 16.5, EMBER, True),
          (" mean cost-to-serve is near-zero. The model works at a ", 16.5, INK),
          ("$50 pledge", 16.5, INK, True), (" as well as a ", 16.5, INK), ("$50,000", 16.5, INK, True),
          (" one — and fiat ramps unlock backers who never held crypto.", 16.5, INK)]], line_spacing=1.32)
footer(s)
notes(s, "Revenue is incentive-aligned: a small fee taken only on milestone release, so Ember earns "
        "only when creators deliver and backers are satisfied. Secondary line: FX/ramp spread on "
        "anchor on/off-ramps. Stress the compliance posture — no token, no token sale, no custody; "
        "pledges are utility. Stellar makes the unit economics work: sub-cent fees mean cost-to-serve "
        "is near-zero, so a $50 pledge is as viable as a $50k one, and fiat ramps expand the top of "
        "funnel to non-crypto backers. If asked about take-rate, keep it directional (low single-digit "
        "%), tiered lower for verified nonprofits/impact.")

# =====================================================================
# SLIDE 14 — WHY EMBER WINS
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "Differentiation")
title(s, "Why Ember wins")
headers = ["", "Funds held\nuntil delivery", "Non-\ncustodial", "Onchain\nproof", "Stablecoin\n-native", "Fiat ramp\n(anchors)"]
data = [
    ("Kickstarter / GoGetFunding", "✗","✗","✗","✗","✗"),
    ("Traditional escrow / banks", "✓","✗","✗","✗","✓"),
    ("Generic crypto crowdfunding", "⚠","✓","✓","⚠","✗"),
    ("Ember on Stellar", "✓","✓","✓","✓","✓"),
]
tx = MX; ty = Inches(1.95)
col0 = Inches(4.0); colw = Inches(1.566); rowh = Inches(0.86)
# header row
for j, h in enumerate(headers):
    x = tx + (col0 if j>0 else Inches(0)) + (Inches(0) if j==0 else (j-1)*colw)
    w = col0 if j==0 else colw
    textbox(s, x, ty, w, Inches(0.75), [[(h, 13, MUTED, True)]], align=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
for i, row in enumerate(data):
    y = ty + Inches(0.8) + i*rowh
    ember_row = (row[0].startswith("Ember"))
    if ember_row:
        hl = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx-Inches(0.15), y-Inches(0.05), CW+Inches(0.0), rowh-Inches(0.06))
        hl.fill.solid(); hl.fill.fore_color.rgb = EMBER_SOFT; hl.line.color.rgb = CARDLINE; hl.line.width=Pt(1); hl.shadow.inherit=False
        try: hl.adjustments[0]=0.2
        except Exception: pass
    textbox(s, tx, y, col0, rowh-Inches(0.1), [[(row[0], 15.5, EMBER if ember_row else INK, ember_row)]], anchor=MSO_ANCHOR.MIDDLE)
    for j in range(1,6):
        x = tx + col0 + (j-1)*colw
        val = row[j]
        col = GREEN if val=="✓" else (EMBER if val=="✗" else MUTED)
        textbox(s, x, y, colw, rowh-Inches(0.1), [[(val, 20, col, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, MX, Inches(6.35), CW, Inches(0.6),
        [[("Milestone escrow ", 16, INK), ("+", 16, EMBER, True), (" weighted onchain voting ", 16, INK), ("+", 16, EMBER, True),
          (" position NFTs ", 16, INK), ("+", 16, EMBER, True), (" native fiat ramps — accountability neither Web2 nor token launchpads offer.", 16, INK)]])
footer(s)
notes(s, "The competitive matrix. Web2 platforms fail every column. Banks/escrow hold funds but are "
        "custodial, offline, and closed. Generic crypto crowdfunding is non-custodial but usually "
        "sells a token at TGE with no milestone gating (the warning marks). Ember is the only row "
        "green across the board — and note the new fifth column versus the old deck: native fiat "
        "ramps via Stellar anchors. That column is only checkable because we're on Stellar; it's the "
        "differentiator that reaches mainstream, non-crypto backers. Deliver the bottom line: the "
        "combination is the moat, not any single feature.")

# =====================================================================
# SLIDE 15 — ROADMAP
# =====================================================================
s = slide(BG)
eyebrow(s, MX, Inches(0.62), "What's next")
title(s, "Roadmap")
phases = [
    ("Now", EMBER, ["Feature-complete product on EVM testnet",
                    "Stellar migration scoped (issue #177)",
                    "Ecosystem integration plan (issue #179)"]),
    ("Next", EMBER, ["Soroban contracts + SEP-50 NFTs on testnet",
                     "Passkey onboarding + SEP-24 peso ramp (Coins.ph)",
                     "Mainnet launch · apply to SCF 7.0"]),
    ("Later", GREEN, ["APAC anchor corridors (VN, ID, IN)",
                      "Quadratic funding & recurring pledges",
                      "Disbursement Platform + RWA/impact rails"]),
]
cw = Inches(3.77); ch = Inches(3.9); gap = Inches(0.26); cy = Inches(2.0)
for i, (h, col, items) in enumerate(phases):
    x = MX + i*(cw+gap)
    card(s, x, cy, cw, ch)
    chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x+Inches(0.3), cy+Inches(0.32), Inches(1.4), Inches(0.5))
    chip.fill.solid(); chip.fill.fore_color.rgb = col; chip.line.fill.background(); chip.shadow.inherit=False
    try: chip.adjustments[0]=0.3
    except Exception: pass
    ctf = chip.text_frame; ctf.paragraphs[0].alignment=PP_ALIGN.CENTER
    _set_run(ctf.paragraphs[0].add_run(), h, 18, WHITE, True)
    textbox(s, x+Inches(0.3), cy+Inches(1.1), cw-Inches(0.6), Inches(2.6),
            [[("•  "+it, 15.5, INK)] for it in items], line_spacing=1.3, space_after=13)
footer(s)
notes(s, "Sequence the plan and show it's already in motion. Now: the product is built and both the "
        "migration (issue #177) and the ecosystem plan (issue #179) are written and scoped. Next: "
        "Soroban contracts and SEP-50 NFTs on testnet, passkey onboarding with the Coins.ph peso "
        "ramp, mainnet, and the SCF 7.0 application. Later: APAC anchor corridors, quadratic funding "
        "and recurring pledges, and the Disbursement Platform / RWA rails. The throughline: near-term "
        "milestones are concrete and low-risk because the product exists; the later items are upside. "
        "This is a fundable, sequenced plan, not a wish list.")

# =====================================================================
# SLIDE 16 — CLOSING
# =====================================================================
s = slide(DARK)
g = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
g.fill.solid(); g.fill.fore_color.rgb = DARK
g.line.fill.background(); g.shadow.inherit=False
acc = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-3.6), Inches(3.4), Inches(9.0), Inches(7.0))
acc.fill.solid(); acc.fill.fore_color.rgb = DARKRED
acc.line.fill.background(); acc.shadow.inherit = False
textbox(s, MX, Inches(1.4), CW, Inches(0.4), [[("EMBER · BUILT ON STELLAR", 13, RGBColor(0xC8,0xB7,0xB2), True)]])
textbox(s, MX, Inches(2.5), CW, Inches(2.0),
        [[("Fund the future,", 58, WHITE, True)],
         [("one milestone at a time", 58, WHITE, True), (".", 58, EMBER_BR, True)]], line_spacing=1.05)
textbox(s, MX, Inches(4.7), CW, Inches(0.6), [[("Accountable crowdfunding — and real impact for Stellar.", 24, CREAM)]])
textbox(s, MX, Inches(5.5), CW, Inches(0.5),
        [[("Strategy & references: GitHub issues #177 (migration) · #179 (ecosystem)", 15, RGBColor(0xC8,0xB7,0xB2))]])
textbox(s, MX, Inches(6.2), CW, Inches(0.5), [[("Thank you — questions welcome.", 17, RGBColor(0xE6,0xE6,0xE6), True)]])
footer(s, "Ember · 2026", dark=True)
notes(s, "Close with the tagline and the dual message: accountable crowdfunding for creators AND "
        "measurable impact for the Stellar ecosystem — stablecoin volume, wallet activations, a new "
        "category. Point to issues #177 and #179 as the full written strategy and reference material. "
        "Then open the floor. Likely Q&A: current chain state (EVM testnet, migrating), why USDC over "
        "USDT (native on Stellar), SCF timing/amount (cite #179 with its verify caveats), and "
        "regulatory posture (non-custodial, no token, KYC offloaded to anchors).")

prs.save("/home/user/ember2/docs/pitch-deck-v3.pptx")
print("Saved docs/pitch-deck-v3.pptx with", len(prs.slides._sldIdLst), "slides")
